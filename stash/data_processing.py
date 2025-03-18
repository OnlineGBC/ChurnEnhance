# file_processing.py

import pandas as pd
import whisper
from pptx import Presentation
import PyPDF2
from docx import Document
from bs4 import BeautifulSoup
from PIL import Image
import pytesseract
from yt_dlp import YoutubeDL
import os

# Whisper model loader
def load_whisper_model(model_name="base"):
    return whisper.load_model(model_name)

# Process Text Files
def process_txt(file_bytes):
    return file_bytes.decode("utf-8")

# Process CSV Files
def process_csv(file_obj):
    df = pd.read_csv(file_obj)
    return df.to_csv(index=False), df

# Process Excel Files
def process_xlsx(file_obj):
    df = pd.read_excel(file_obj, engine="openpyxl")
    return df.to_csv(index=False), df

# Process PPTX Files
def process_pptx(file_obj):
    prs = Presentation(file_obj)
    ppt_text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                ppt_text.append(shape.text)
    return "\n".join(ppt_text)

# Process PDF Files
def process_pdf(file_obj):
    pdf_reader = PyPDF2.PdfReader(file_obj)
    text_parts = []
    for page in pdf_reader.pages:
        text_parts.append(page.extract_text() or "")
    return "".join(text_parts)

# Process DOCX Files
def process_docx(file_obj):
    doc = Document(file_obj)
    doc_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            doc_text.append(para.text)
    return "\n".join(doc_text)

# Process MSG Files
def process_msg(file_obj):
    import extract_msg
    msg = extract_msg.Message(file_obj)
    return msg.body

# Process HTML Files
def process_html(file_obj):
    html = file_obj.read().decode("utf-8")
    soup = BeautifulSoup(html, "html.parser")
    return soup.get_text(separator="\n")

# Process Images (OCR)
def process_image(file_obj):
    image = Image.open(file_obj)
    return pytesseract.image_to_string(image)

# Process Audio/Video (Whisper)
def transcribe_audio_video(whisper_model, file_path):
    result = whisper_model.transcribe(file_path)
    return result.get("text", "")

# Media Downloading Functionality
def download_media(url: str, name: str, format_choice: str, output_folder: str) -> tuple:
    ydl_opts = {
        'outtmpl': f'{os.path.join(output_folder, name)}.%(ext)s',
        'quiet': True,
        'no_warnings': True
    }

    if format_choice == "MP3 (Audio-only)":
        ydl_opts.update({
            'format': 'bestaudio/best',
            'postprocessors': [{
                'key': 'FFmpegExtractAudio',
                'preferredcodec': 'mp3',
                'preferredquality': '192',
            }],
        })
    else:  # MP4 (Video)
        ydl_opts['format'] = 'bestvideo+bestaudio/best'

    try:
        with YoutubeDL(ydl_opts) as ydl:
            ydl.download([url])
        return True, "Download completed successfully."
    except Exception as e:
        return False, f"Download failed: {e}"
